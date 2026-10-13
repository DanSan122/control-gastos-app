import csv
from io import BytesIO, StringIO
from datetime import datetime
from decimal import Decimal
from unittest import skipUnless
from django.contrib.auth import get_user_model
from django.db import connection
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient
from apps.accounts.models import Account
from apps.categories.models import Category, Subcategory
from apps.transactions.models import Transaction


class ReportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            "report@example.test", "Bosque!Prueba2026"
        )
        self.other = get_user_model().objects.create_user(
            "report-other@example.test", "Bosque!Prueba2026"
        )
        self.client = APIClient(enforce_csrf_checks=True)
        self.client.force_login(self.user)
        self.account = Account.objects.create(
            owner=self.user, name="Banco reportes", kind="BANK"
        )
        self.category = Category.objects.filter(
            owner=self.user, kind="EXPENSE"
        ).first()
        self.subcategory = Subcategory.objects.filter(
            category=self.category
        ).first()
        self.query = "period=custom&start=2026-01-01&end=2026-01-31"

    def movement(self, amount="10.25", **extra):
        return Transaction.objects.create(
            **{
                **dict(
                    owner=self.user,
                    account=self.account,
                    category=self.category,
                    area=self.category.area,
                    subcategory=self.subcategory,
                    kind="EXPENSE",
                    amount=Decimal(amount),
                    occurred_at=datetime.fromisoformat(
                        "2026-01-10T12:00:00-05:00"
                    ),
                    payment_method="YAPE",
                    description="Compra prueba",
                ),
                **extra,
            }
        )

    def get(self, path="preview", query=None):
        return self.client.get(
            f"/api/v1/reports/{path}/?{query or self.query}"
        )

    def test_report_exact_aggregate_status_transfer_credit_and_filters(self):
        expense = self.movement()
        card = Account.objects.create(
            owner=self.user, name="Tarjeta", kind="CREDIT"
        )
        self.movement("20.00", account=card)
        self.movement("999.00", status="VOID")
        income = Category.objects.filter(
            owner=self.user, kind="INCOME"
        ).first()
        self.movement(
            "100",
            kind="INCOME",
            category=income,
            area=income.area,
            subcategory=None,
        )
        self.movement(
            "20",
            kind="TRANSFER",
            category=None,
            area=None,
            subcategory=None,
            destination_account=card,
        )
        data = self.get().json()
        self.assertEqual(
            (data["income"], data["expense"], data["net_flow"]),
            ("100.00", "30.25", "69.75"),
        )
        self.assertEqual(data["movements"]["count"], 4)
        self.assertEqual(
            self.get(query=self.query + "&status=").json()["movements"][
                "count"
            ],
            4,
        )
        query = f"{self.query}&account={self.account.pk}&area={self.category.area_id}&category={self.category.pk}&subcategory={self.subcategory.pk}&kind=EXPENSE&payment_method=YAPE&status=POSTED"
        data = self.get(query=query).json()
        self.assertEqual(data["expense"], "10.25")
        self.assertEqual(
            data["movements"]["results"][0]["id"], str(expense.pk)
        )
        data = self.get(query=self.query + "&status=VOID").json()
        self.assertEqual(data["expense"], "0.00")
        self.assertEqual(data["movements"]["count"], 1)

    def test_csv_bom_iso_decimal_and_formula_guard(self):
        self.movement(description=' \t=HYPERLINK("evil")', reference="@SUM(1)")
        response = self.get("export/csv")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(rows[1][9], "10.25")
        self.assertTrue(rows[1][12].startswith("'"))
        self.assertTrue(rows[1][13].startswith("'"))
        self.assertEqual(rows[1][1], "2026-01-10T12:00:00-05:00")
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertIn("text/csv", response["Content-Type"])
        self.assertEqual(response["Cache-Control"], "private, no-store")

    def test_xlsx_five_sheets_exact_strings_styles_and_no_formulas(self):
        self.movement(description="=1+1", reference="+cmd")
        response = self.get("export/xlsx")
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(
            workbook.sheetnames,
            [
                "Resumen financiero",
                "Movimientos",
                "Gastos por área",
                "Gastos por categoría",
                "Evolución temporal",
            ],
        )
        self.assertEqual(workbook["Movimientos"]["J5"].value, "10.25")
        self.assertEqual(workbook["Movimientos"]["J5"].number_format, '"S/ "@')
        self.assertEqual(workbook["Movimientos"]["M5"].value, "=1+1")
        self.assertEqual(workbook["Resumen financiero"]["B6"].value, "10.25")
        self.assertTrue(workbook["Movimientos"].auto_filter.ref)
        self.assertEqual(workbook["Movimientos"].freeze_panes, "A5")
        for sheet in workbook:
            for row in sheet:
                self.assertTrue(all(cell.data_type != "f" for cell in row))

    @skipUnless(
        connection.vendor == "postgresql",
        "NUMERIC extremo requiere PostgreSQL",
    )
    def test_xlsx_preserves_maximum_money_precision(self):
        self.movement("9999999999999999.99")
        response = self.get("export/xlsx")
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(
            book["Movimientos"]["J5"].value, "9999999999999999.99"
        )

    def test_pdf_and_empty_documents(self):
        response = self.get("export/pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertIn(b"%%EOF", response.content)
        self.assertEqual(self.get().json()["expense"], "0.00")
        self.assertEqual(
            load_workbook(BytesIO(self.get("export/xlsx").content))[
                "Movimientos"
            ].max_row,
            4,
        )
        self.assertEqual(
            len(
                list(
                    csv.reader(
                        StringIO(
                            self.get("export/csv").content.decode("utf-8-sig")
                        )
                    )
                )
            ),
            1,
        )

    def test_pdf_long_safe_text_and_many_pages(self):
        for index in range(35):
            self.movement(
                description=("<script> & texto largo " * 22)[:500],
                reference=f"Ref {index}",
            )
        response = self.get("export/pdf")
        self.assertEqual(response.status_code, 200)
        self.assertGreater(response.content.count(b"/Type /Page"), 3)

    def test_owner_and_private_permissions(self):
        account = Account.objects.create(owner=self.other, name="Ajena")
        category = Category.objects.filter(
            owner=self.other, kind="EXPENSE"
        ).first()
        self.movement(
            "777",
            owner=self.other,
            account=account,
            area=category.area,
            category=category,
            subcategory=None,
        )
        self.assertEqual(self.get().json()["movements"]["count"], 0)
        self.assertEqual(
            self.get(query=self.query + f"&account={account.pk}").status_code,
            400,
        )
        self.assertNotIn(b"777", self.get("export/csv").content)
        self.client.logout()
        for path in ("preview", "export/xlsx", "export/csv", "export/pdf"):
            self.assertEqual(self.get(path).status_code, 403)

    @override_settings(REPORT_EXPORT_ROWS=1, REPORT_PDF_ROWS=1)
    def test_ranges_volume_and_invalid_filters(self):
        self.movement()
        self.movement()
        for kind in ("xlsx", "csv", "pdf"):
            self.assertEqual(self.get(f"export/{kind}").status_code, 400)
        for query in (
            "period=custom&start=2024-01-01&end=2026-01-01",
            "account=invalid",
            "currency=USD",
            "status=wrong",
            "payment_method=wrong",
            "granularity=wrong",
            "owner=bad",
        ):
            self.assertEqual(self.get(query=query).status_code, 400)

    def test_pagination_calendar_boundaries_and_unspecified_method(self):
        self.movement(
            "1",
            occurred_at=datetime.fromisoformat("2026-01-01T04:59:59+00:00"),
        )
        self.movement(
            "2",
            occurred_at=datetime.fromisoformat("2026-01-01T05:00:00+00:00"),
            payment_method="",
        )
        self.movement(
            "4",
            occurred_at=datetime.fromisoformat("2026-02-01T05:00:00+00:00"),
        )
        data = self.get(
            query=self.query + "&payment_method=UNSPECIFIED&page_size=1"
        ).json()
        self.assertEqual(data["expense"], "2.00")
        self.assertEqual(data["movements"]["count"], 1)

    @override_settings(REPORT_PREVIEW_LIMIT=1, REPORT_EXPORT_LIMIT=1)
    def test_atomic_shared_throttle(self):
        self.assertEqual(self.get().status_code, 200)
        response = self.get()
        self.assertEqual(response.status_code, 429)
        self.assertIn("Retry-After", response)
        self.assertEqual(self.get("export/csv").status_code, 200)
        self.assertEqual(self.get("export/csv").status_code, 429)
